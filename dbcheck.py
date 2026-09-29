# import sqlite3
# import json
# from pathlib import Path
# import pandas as pd
# from datetime import datetime
# import openpyxl
# from openpyxl.utils import get_column_letter

# DB_PATH = "data/raw_repository.db"
# OUTPUT_DIR = Path("output")
# OUTPUT_DIR.mkdir(exist_ok=True)

# def main(limit=10):
#     if not Path(DB_PATH).exists():
#         print(f"Database not found at {DB_PATH}")
#         return

#     conn = sqlite3.connect(DB_PATH)
#     conn.row_factory = sqlite3.Row
#     cursor = conn.cursor()

#     # Join both tables and select all columns, alias to avoid conflicts
#     cursor.execute(f"""
#         SELECT
#             s.id AS staging_id,
#             s.metadata_hash,
#             s.source_domain,
#             s.source_url,
#             s.title,
#             s.author,
#             s.published_utc,
#             s.raw_html,
#             s.original_source_url,
#             s.source_status,
#             s.retrieved_utc,
#             s.status AS staging_status,
#             s.processing_error AS staging_error,
#             p.id AS processed_id,
#             p.summary,
#             p.iocs_json,
#             p.threat_metadata_json,
#             p.mitre_techniques_json,
#             p.llm_raw_response,
#             p.processing_status,
#             p.processing_error,
#             p.processing_timestamp,
#             p.confidence_score,
#             p.verification_status,
#             p.hallucination_count,
#             p.evidence_json,
#             p.target_organization
#         FROM staging_raw s
#         JOIN processed_articles p ON s.metadata_hash = p.metadata_hash
#         ORDER BY p.processing_timestamp DESC
#         LIMIT ?
#     """, (limit,))

#     rows = cursor.fetchall()
#     if not rows:
#         print("No processed articles found.")
#         return

#     print(f"\n{'='*80}")
#     print(f"Exporting {len(rows)} processed articles to Excel...")
#     print('='*80)

#     # Convert rows to list of dicts, with JSON fields prettified
#     data = []
#     for row in rows:
#         record = dict(row)
#         # Pretty‑print JSON fields
#         for key in ['iocs_json', 'threat_metadata_json', 'mitre_techniques_json', 'evidence_json']:
#             if key in record and record[key]:
#                 try:
#                     parsed = json.loads(record[key])
#                     record[key] = json.dumps(parsed, indent=2, ensure_ascii=False)
#                 except:
#                     pass  # keep as is
#         # Truncate raw_html and llm_raw_response to avoid huge cells
#         if record.get('raw_html'):
#             record['raw_html'] = record['raw_html'][:500] + '...' if len(record['raw_html']) > 500 else record['raw_html']
#         if record.get('llm_raw_response'):
#             record['llm_raw_response'] = record['llm_raw_response'][:500] + '...' if len(record['llm_raw_response']) > 500 else record['llm_raw_response']
#         data.append(record)

#     df = pd.DataFrame(data)

#     timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
#     excel_path = OUTPUT_DIR / f"processed_articles_full_{timestamp}.xlsx"

#     # Write to Excel with auto‑adjusted column widths
#     with pd.ExcelWriter(excel_path, engine='openpyxl') as writer:
#         df.to_excel(writer, sheet_name="Processed", index=False)
#         workbook = writer.book
#         worksheet = writer.sheets["Processed"]
#         for column in worksheet.columns:
#             max_length = 0
#             column_letter = get_column_letter(column[0].column)
#             for cell in column:
#                 try:
#                     if cell.value:
#                         max_length = max(max_length, len(str(cell.value)))
#                 except:
#                     pass
#             adjusted_width = min(max_length + 2, 60)
#             worksheet.column_dimensions[column_letter].width = adjusted_width

#     print(f"\n✅ Excel file saved to: {excel_path}")
#     conn.close()

# if __name__ == "__main__":
#     main(limit=10)
import sqlite3
import json
from pathlib import Path
import pandas as pd
from datetime import datetime
import openpyxl
from openpyxl.utils import get_column_letter
from openpyxl.styles import Alignment  # <-- NEW: Required for rendering newlines

DB_PATH = "data/raw_repository.db"
OUTPUT_DIR = Path("output")
OUTPUT_DIR.mkdir(exist_ok=True)

def main(limit=10):
    if not Path(DB_PATH).exists():
        print(f"Database not found at {DB_PATH}")
        return

    conn = sqlite3.connect(DB_PATH)
    conn.row_factory = sqlite3.Row
    cursor = conn.cursor()

    # Join both tables and select all columns, alias to avoid conflicts
    cursor.execute(f"""
        SELECT
            s.id AS staging_id,
            s.metadata_hash,
            s.source_domain,
            s.source_url,
            s.title,
            s.author,
            s.published_utc,
            s.raw_html,
            s.original_source_url,
            s.source_status,
            s.retrieved_utc,
            s.status AS staging_status,
            s.processing_error AS staging_error,
            p.id AS processed_id,
            p.summary,
            p.iocs_json,
            p.threat_metadata_json,
            p.mitre_techniques_json,
            p.llm_raw_response,
            p.processing_status,
            p.processing_error,
            p.processing_timestamp,
            p.confidence_score,
            p.verification_status,
            p.hallucination_count,
            p.evidence_json,
            p.target_organization
        FROM staging_raw s
        JOIN processed_articles p ON s.metadata_hash = p.metadata_hash
        ORDER BY p.processing_timestamp DESC
        LIMIT ?
    """, (limit,))

    rows = cursor.fetchall()
    if not rows:
        print("No processed articles found.")
        return

    print(f"\n{'='*80}")
    print(f"Exporting {len(rows)} processed articles to Excel...")
    print('='*80)

    data = []
    for row in rows:
        record = dict(row)
        # --- FIX: Added target_organization to the list ---
        for key in ['iocs_json', 'threat_metadata_json', 'mitre_techniques_json', 'evidence_json', 'target_organization']:
            if key in record and record[key]:
                try:
                    parsed = json.loads(record[key])
                    record[key] = json.dumps(parsed, indent=2, ensure_ascii=False)
                except:
                    pass  # keep as is
        
        # Truncate raw_html and llm_raw_response to avoid huge cells
        if record.get('raw_html'):
            record['raw_html'] = record['raw_html'][:500] + '...' if len(record['raw_html']) > 500 else record['raw_html']
        if record.get('llm_raw_response'):
            record['llm_raw_response'] = record['llm_raw_response'][:500] + '...' if len(record['llm_raw_response']) > 500 else record['llm_raw_response']
        data.append(record)

    df = pd.DataFrame(data)

    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    excel_path = OUTPUT_DIR / f"processed_articles_full_{timestamp}.xlsx"

    # Write to Excel with auto‑adjusted column widths and text wrapping
    with pd.ExcelWriter(excel_path, engine='openpyxl') as writer:
        df.to_excel(writer, sheet_name="Processed", index=False)
        worksheet = writer.sheets["Processed"]
        
        # --- FIX: Alignment styling for proper JSON rendering ---
        wrap_alignment = Alignment(wrap_text=True, vertical='top')

        for column in worksheet.columns:
            max_length = 0
            column_letter = get_column_letter(column[0].column)
            for cell in column:
                try:
                    if cell.value:
                        cell.alignment = wrap_alignment
                        # Measure longest single line, not total character count, to prevent 
                        # columns from becoming extremely wide when text wraps.
                        lines = str(cell.value).split('\n')
                        longest_line = max([len(line) for line in lines])
                        max_length = max(max_length, longest_line)
                except:
                    pass
            adjusted_width = min(max_length + 2, 70)  # Max width cap
            worksheet.column_dimensions[column_letter].width = adjusted_width

    print(f"\n✅ Excel file saved to: {excel_path}")
    conn.close()

if __name__ == "__main__":
    main(limit=10)